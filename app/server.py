#!/usr/bin/env python3
import json
import os
import re
import sys
import traceback
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path

import fitz  # PyMuPDF
import openpyxl
import torch
from docx import Document
from PIL import Image
from qwen_vl_utils import process_vision_info
from transformers import AutoProcessor, Qwen2_5_VLForConditionalGeneration

MODEL_PATH = os.environ.get("MODEL_PATH", "/models/Qwen2.5-VL-7B-Instruct")
INDEX_DIR = Path(os.environ.get("MC3_INDEX_DIR", "/app/index"))


class RAGServer:

    def __init__(self):
        print(
            f"[Server] Loading Qwen2.5-VL-7B-Instruct from {MODEL_PATH} onto AMD GPU...",
            flush=True,
        )
        self.device = "cuda" if torch.cuda.is_available() else "cpu"

        # Load weights on AMD ROCm GPU
        self.model = Qwen2_5_VLForConditionalGeneration.from_pretrained(
            MODEL_PATH,
            torch_dtype=torch.bfloat16 if self.device == "cuda" else torch.float32,
            device_map="auto" if self.device == "cuda" else None,
        )
        self.processor = AutoProcessor.from_pretrained(MODEL_PATH)
        self.docs = []

        # Load existing index if already built
        index_file = INDEX_DIR / "corpus_index.json"
        if index_file.exists():
            try:
                self.docs = json.loads(index_file.read_text(encoding="utf-8"))
                print(f"[Server] Restored {len(self.docs)} cached documents.", flush=True)
            except Exception:
                pass

        print("[Server] Engine loaded and ready!", flush=True)

    def ask_ai(self, messages: list, max_tokens: int = 512) -> str:
        text = self.processor.apply_chat_template(
            messages, tokenize=False, add_generation_prompt=True
        )
        image_inputs, video_inputs = process_vision_info(messages)
        inputs = self.processor(
            text=[text],
            images=image_inputs,
            videos=video_inputs,
            padding=True,
            return_tensors="pt",
        ).to(self.device)

        with torch.no_grad():
            generated_ids = self.model.generate(
                **inputs, max_new_tokens=max_tokens, do_sample=False
            )
            trimmed_ids = [
                out_ids[len(in_ids) :]
                for in_ids, out_ids in zip(inputs.input_ids, generated_ids)
            ]
            response = self.processor.batch_decode(
                trimmed_ids,
                skip_special_tokens=True,
                clean_up_tokenization_spaces=False,
            )[0]
        return response.strip()

    def extract_with_vlm(
        self, image_path: str, doc_type: str = "document page"
    ) -> str:
        messages = [
            {
                "role": "user",
                "content": [
                    {"type": "image", "image": f"file://{image_path}"},
                    {
                        "type": "text",
                        "text": (
                            f"You are an expert technical document reader. Extract ALL text, "
                            f"tables (in markdown table format), pinouts, specifications, "
                            f"and diagram labels visible in this {doc_type}. Do not summarize; "
                            f"preserve exact numbers, units, and codes verbatim."
                        ),
                    },
                ],
            }
        ]
        return self.ask_ai(messages, max_tokens=1024)

    def extract_text_from_file(self, file_path: Path) -> str:
        suffix = file_path.suffix.lower()

        # 1. Plain text / Code / Logs / CSV
        if suffix in [".txt", ".log", ".py", ".csv"]:
            return file_path.read_text(encoding="utf-8", errors="ignore")

        # 2. Images (.png, .jpg)
        if suffix in [".png", ".jpg", ".jpeg"]:
            with Image.open(file_path) as img:
                img.verify()
            return self.extract_with_vlm(
                str(file_path.resolve()), doc_type="diagram/image"
            )

        # 3. PDF files
        if suffix == ".pdf":
            doc = fitz.open(str(file_path))
            if doc.is_encrypted:
                print(
                    f"[Index] Skipping encrypted PDF: {file_path.name}", flush=True
                )
                doc.close()
                return ""

            extracted_pages = []
            temp_img_path = "/tmp/pdf_page_render.png"
            for page_num in range(len(doc)):
                page = doc[page_num]
                digital_text = page.get_text("text").strip()

                # Render page to check for visual tables / diagrams
                pix = page.get_pixmap(dpi=150)
                pix.save(temp_img_path)
                vlm_markdown = self.extract_with_vlm(
                    temp_img_path, doc_type=f"PDF page {page_num + 1}"
                )

                extracted_pages.append(
                    f"--- Page {page_num + 1} ---\n{vlm_markdown}\n{digital_text}"
                )
            doc.close()
            return "\n\n".join(extracted_pages)

        # 4. Word Documents (.docx)
        if suffix == ".docx":
            doc = Document(str(file_path))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for t_idx, table in enumerate(doc.tables):
                parts.append(f"\n[Table {t_idx + 1}]")
                for row in table.rows:
                    row_text = " | ".join(
                        cell.text.strip().replace("\n", " ")
                        for cell in row.cells
                    )
                    if row_text.strip():
                        parts.append(f"| {row_text} |")
            return "\n".join(parts)

        # 5. Excel Spreadsheets (.xlsx)
        if suffix == ".xlsx":
            wb = openpyxl.load_workbook(str(file_path), data_only=True)
            parts = []
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                parts.append(f"### Sheet: {sheet}")
                for row in ws.iter_rows(values_only=True):
                    vals = [
                        str(v).strip() if v is not None else "" for v in row
                    ]
                    if any(vals):
                        parts.append("| " + " | ".join(vals) + " |")
            return "\n".join(parts)

        return ""

    def build_index(self, corpus_str: str) -> dict:
        corpus = Path(corpus_str)
        INDEX_DIR.mkdir(parents=True, exist_ok=True)
        self.docs = []
        valid_exts = {
            ".pdf",
            ".docx",
            ".xlsx",
            ".csv",
            ".txt",
            ".log",
            ".py",
            ".png",
            ".jpg",
            ".jpeg",
        }

        for file_path in sorted(corpus.rglob("*")):
            if file_path.is_dir():
                continue

            rel_path = str(file_path.relative_to(corpus)).replace("\\", "/")
            suffix = file_path.suffix.lower()

            if suffix not in valid_exts:
                print(
                    f"[Index] Skipping unknown extension: {rel_path}", flush=True
                )
                continue

            try:
                # Test read permissions (skips chmod 000 unreadables)
                with open(file_path, "rb") as f:
                    f.read(16)

                text = self.extract_text_from_file(file_path)
                if text and text.strip():
                    self.docs.append({"path": rel_path, "text": text.strip()})
                    print(f"[Index] Indexed {rel_path}", flush=True)
            except (PermissionError, OSError) as e:
                print(
                    f"[Index] Skipping unreadable file {rel_path}: {e}",
                    flush=True,
                )
            except Exception as e:
                print(f"[Index] Error parsing {rel_path}: {e}", flush=True)

        (INDEX_DIR / "corpus_index.json").write_text(
            json.dumps(self.docs, indent=2), encoding="utf-8"
        )
        return {"status": "ok", "count": len(self.docs)}

    def answer_question(self, corpus_str: str, query: str) -> dict:
        if not self.docs:
            index_file = INDEX_DIR / "corpus_index.json"
            if index_file.exists():
                self.docs = json.loads(index_file.read_text(encoding="utf-8"))
            else:
                self.build_index(corpus_str)

        # Do not cite withdrawn documents when valid revisions exist
        active_docs = [
            d for d in self.docs if "withdrawn" not in d["path"].lower()
        ]

        context_blocks = []
        for i, doc in enumerate(active_docs):
            snippet = doc["text"][:7000]
            context_blocks.append(
                f"=== DOCUMENT PATH: {doc['path']} ===\n{snippet}"
            )

        full_context = "\n\n".join(context_blocks)

        prompt = f"""You are a precise technical question-answering system. Answer the question using ONLY the provided documents.

RULES:
1. Return ONLY a valid JSON object with keys: "answer", "citations", "confidence".
2. "answer": The exact value only (e.g. "94", "Q3 FY27", "REV-C2"). NO full sentences, NO conversational filler.
3. "citations": A list of relative paths (e.g. ["specs/tq40_datasheet_r2.pdf"]).
   - Cite a file ONLY if removing it makes the answer impossible to determine.
   - For multi-hop questions (e.g. linking a log error to a bug entry), cite BOTH files.
   - Never cite withdrawn documents.
4. If the answer cannot be determined from the documents, return:
   {{"answer": "", "citations": [], "confidence": 0.0}}

DOCUMENTS:
{full_context}

QUESTION: {query}

JSON Output:"""

        messages = [
            {"role": "user", "content": [{"type": "text", "text": prompt}]}
        ]
        raw_output = self.ask_ai(messages, max_tokens=100)

        match = re.search(r"\{.*\}", raw_output, re.DOTALL)
        if not match:
            return {"answer": "", "citations": [], "confidence": 0.0}

        try:
            parsed = json.loads(match.group(0))
            ans = str(parsed.get("answer", "")).strip()
            cites = [str(c).strip() for c in parsed.get("citations", []) if c]

            if not ans or ans.lower() in [
                "none",
                "null",
                "n/a",
                "unknown",
                "i don't know",
            ]:
                return {"answer": "", "citations": [], "confidence": 0.0}

            valid_paths = {d["path"] for d in active_docs}
            cites = [c for c in cites if c in valid_paths]
            if not cites:
                return {"answer": "", "citations": [], "confidence": 0.0}

            return {
                "answer": ans,
                "citations": cites,
                "confidence": float(parsed.get("confidence", 0.9)),
            }
        except Exception:
            return {"answer": "", "citations": [], "confidence": 0.0}


rag_engine = None


class Handler(BaseHTTPRequestHandler):

    def do_GET(self):
        if self.path == "/health":
            self.send_response(200)
            self.end_headers()
            self.wfile.write(b"OK")

    def do_POST(self):
        length = int(self.headers.get("Content-Length", 0))
        body = json.loads(self.rfile.read(length).decode("utf-8"))
        try:
            if self.path == "/index":
                res = rag_engine.build_index(body["corpus"])
            elif self.path == "/query":
                res = rag_engine.answer_question(body["corpus"], body["query"])
            else:
                res = {"error": "unknown endpoint"}
            self.send_response(200)
        except Exception as e:
            traceback.print_exc()
            res = {
                "answer": "",
                "citations": [],
                "confidence": 0.0,
                "error": str(e),
            }
            self.send_response(500)

        self.send_header("Content-Type", "application/json")
        self.end_headers()
        self.wfile.write(json.dumps(res).encode("utf-8"))

    def log_message(self, format, *args):
        # Silence default HTTP access logs to keep terminal output clean
        return


if __name__ == "__main__":
    rag_engine = RAGServer()
    server = HTTPServer(("127.0.0.1", 8000), Handler)
    print("[Server] Listening on http://127.0.0.1:8000", flush=True)
    server.serve_forever()
